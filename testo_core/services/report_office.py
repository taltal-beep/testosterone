"""The manager report as an Excel workbook and a PowerPoint deck.

Both use native Office objects, not pictures: workbook tables people can sort and
filter, and slides whose charts are real charts (edit the data, restyle them to a
company template). They need the ``export`` extra (``openpyxl``, ``python-pptx``).
The numbers and wording come from :mod:`testo_core.services.manager_report`, the
same as the ``/report`` page.
"""

from __future__ import annotations

import io
from datetime import UTC, datetime
from typing import Any

from testo_core.services.manager_report import ManagerReport, NamedTest
from testo_core.services.report_export import CYCLE_COLUMNS, DEFAULT_BRAND, ExportUnavailable

_HINT = "pip install 'testo-core[export]'"

# The report page's palette (frontend/src/features/report/palette.ts).
INK = "101828"
MUTED = "475467"
RULE = "E4E7EC"
WASH = "F2F4F7"
BRAND = "1D4ED8"
GOOD = "067647"
BAD = "B42318"
WARN = "B54708"


def _date(ts: float) -> str:
    when = datetime.fromtimestamp(ts, tz=UTC)
    return f"{when:%b} {when.day}, {when:%Y}"


def period_range(report: ManagerReport) -> str:
    if report.days <= 1:
        return _date(report.end)
    return f"{_date(report.start)} – {_date(report.end)}"


def scope_text(report: ManagerReport) -> str:
    return ", ".join(report.scope) if report.scope else "All cycles"


def _pct(v: float | None) -> str:
    return "–" if v is None else f"{v:.1f}%"


def _delta(now: float | None, before: float | None) -> str:
    if now is None or before is None:
        return ""
    d = now - before
    if abs(d) < 0.05:
        return "no change"
    return f"{'+' if d > 0 else '−'}{abs(d):.1f} pts"


def kpi_cards(report: ManagerReport) -> list[tuple[str, str, str, str]]:
    """(label, value, note, colour) for the four headline numbers, as on the page."""
    k = report.kpis
    score = k.health_score
    score_colour = MUTED if score is None else GOOD if score >= 80 else WARN if score >= 60 else BAD
    green_pct = round(100 * k.green / k.runs) if k.runs else None
    prev_green = round(100 * k.prev_green / k.prev_runs) if k.prev_runs else None
    green_delta = _delta(green_pct, prev_green)
    cards = [
        (
            "Test health",
            "–" if score is None else str(score),
            f"Grade {k.health_grade} · out of 100" if k.health_grade else "No runs scored yet",
            score_colour,
        ),
        (
            "Tests passing",
            _pct(k.pass_rate),
            _delta(k.pass_rate, k.prev_pass_rate) or "No earlier period",
            INK,
        ),
        (
            "Green builds",
            "–" if green_pct is None else f"{green_pct}%",
            f"{k.green} of {k.runs} runs" + (f" · {green_delta}" if green_delta else ""),
            INK,
        ),
    ]
    if k.coverage_pct is not None:
        cards.append(
            (
                "Line coverage",
                _pct(k.coverage_pct),
                _delta(k.coverage_pct, k.prev_coverage_pct) or "No earlier period",
                INK,
            )
        )
    else:
        attention = k.failing_tests + k.newly_flaky
        cards.append(
            (
                "Needs attention",
                str(attention),
                f"{k.failing_tests} failing · {k.newly_flaky} newly flaky",
                BAD if attention else GOOD,
            )
        )
    return cards


def _lists(report: ManagerReport) -> list[tuple[str, list[NamedTest]]]:
    return [
        ("Still failing", report.still_failing),
        ("Turned flaky", report.newly_flaky),
        ("Fixed", report.fixed),
        ("Got slower", report.slower),
    ]


# -- Excel ----------------------------------------------------------------------------


def render_xlsx(report: ManagerReport, brand: str = DEFAULT_BRAND) -> bytes:
    """A workbook: Summary, Trend (with a chart), Cycles, Tests and Findings sheets."""
    try:
        from openpyxl import Workbook
        from openpyxl.chart import LineChart, Reference
        from openpyxl.styles import Alignment, Font, PatternFill
        from openpyxl.utils import get_column_letter
    except ModuleNotFoundError:
        raise ExportUnavailable(f"Excel export needs openpyxl: {_HINT}") from None

    head_font = Font(bold=True, color="FFFFFF")
    head_fill = PatternFill("solid", fgColor=BRAND)

    def table(ws: Any, headers: list[str], rows: list[list[Any]], start: int = 1) -> None:
        for col, text in enumerate(headers, 1):
            cell = ws.cell(row=start, column=col, value=text)
            cell.font, cell.fill = head_font, head_fill
        for r, values in enumerate(rows, start + 1):
            for col, value in enumerate(values, 1):
                ws.cell(row=r, column=col, value=value)
        for col, text in enumerate(headers, 1):
            longest = max([len(str(text))] + [len(str(v[col - 1] or "")) for v in rows])
            ws.column_dimensions[get_column_letter(col)].width = min(60, max(10, longest + 2))
        ws.freeze_panes = ws.cell(row=start + 1, column=1)
        if rows:
            last = get_column_letter(len(headers))
            ws.auto_filter.ref = f"A{start}:{last}{start + len(rows)}"

    wb = Workbook()
    ws = wb.active
    ws.title = "Summary"
    ws["A1"] = report.title
    ws["A1"].font = Font(bold=True, size=16, color=INK)
    ws["A2"] = f"{brand} · {period_range(report)} · {scope_text(report)}"
    ws["A2"].font = Font(color=MUTED)
    row = 4
    for label, value, note, _ in kpi_cards(report):
        ws.cell(row=row, column=1, value=label).font = Font(bold=True)
        ws.cell(row=row, column=2, value=value)
        ws.cell(row=row, column=3, value=note).font = Font(color=MUTED)
        row += 1
    row += 1
    ws.cell(row=row, column=1, value="Summary").font = Font(bold=True)
    for line in report.summary:
        row += 1
        ws.cell(row=row, column=1, value=line).alignment = Alignment(wrap_text=False)
    ws.column_dimensions["A"].width = 22
    ws.column_dimensions["B"].width = 14
    ws.column_dimensions["C"].width = 40

    trend = wb.create_sheet("Trend")
    table(
        trend,
        ["Day", "Runs", "Green runs", "Pass rate %"],
        [[p.day, p.runs, p.green, p.pass_rate] for p in report.trend],
    )
    if report.trend:
        chart = LineChart()
        chart.title = "Pass rate by day"
        chart.y_axis.title = "%"
        chart.height, chart.width = 8, 18
        chart.add_data(
            Reference(trend, min_col=4, min_row=1, max_row=len(report.trend) + 1),
            titles_from_data=True,
        )
        chart.set_categories(Reference(trend, min_col=1, min_row=2, max_row=len(report.trend) + 1))
        trend.add_chart(chart, "F2")

    cycles = wb.create_sheet("Cycles")
    table(
        cycles,
        [label for _, label in CYCLE_COLUMNS],
        [[getattr(c, key) for key, _ in CYCLE_COLUMNS] for c in report.cycles],
    )

    tests = wb.create_sheet("Tests")
    table(
        tests,
        ["List", "Cycle", "Test", "Change"],
        [[name, t.cycle, t.name, t.detail] for name, items in _lists(report) for t in items],
    )

    findings = wb.create_sheet("Findings")
    table(
        findings,
        ["Cycle", "Finding", "Detail"],
        [[f.cycle, f.title, f.detail] for f in report.findings],
    )

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# -- PowerPoint -------------------------------------------------------------------------


def render_pptx(report: ManagerReport, brand: str = DEFAULT_BRAND) -> bytes:
    """A 16:9 deck: title, headline numbers, trend charts, cycles, what needs attention, fixes."""
    try:
        from pptx import Presentation
        from pptx.chart.data import CategoryChartData
        from pptx.dml.color import RGBColor
        from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION
        from pptx.enum.text import PP_ALIGN
        from pptx.util import Emu, Inches, Pt
    except ModuleNotFoundError:
        raise ExportUnavailable(f"PowerPoint export needs python-pptx: {_HINT}") from None

    def rgb(hex_: str) -> Any:
        return RGBColor.from_string(hex_)

    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    blank = prs.slide_layouts[6]
    width = prs.slide_width

    def text(
        slide: Any,
        x: float,
        y: float,
        w: float,
        h: float,
        value: str,
        *,
        size: int = 14,
        bold: bool = False,
        colour: str = INK,
        align: Any = None,
    ) -> Any:
        box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
        frame = box.text_frame
        frame.word_wrap = True
        para = frame.paragraphs[0]
        para.text = value
        para.font.size, para.font.bold, para.font.color.rgb = Pt(size), bold, rgb(colour)
        if align is not None:
            para.alignment = align
        return box

    def slide_with_title(title: str) -> Any:
        slide = prs.slides.add_slide(blank)
        bar = slide.shapes.add_shape(1, 0, 0, width, Inches(0.12))
        bar.fill.solid()
        bar.fill.fore_color.rgb = rgb(BRAND)
        bar.line.fill.background()
        text(slide, 0.6, 0.35, 10, 0.7, title, size=28, bold=True)
        text(
            slide,
            9.0,
            0.45,
            3.8,
            0.5,
            f"{brand} · {period_range(report)}",
            size=11,
            colour=MUTED,
            align=PP_ALIGN.RIGHT,
        )
        return slide

    def bullets(
        slide: Any, x: float, y: float, w: float, h: float, lines: list[str], size: int = 16
    ) -> None:
        box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
        frame = box.text_frame
        frame.word_wrap = True
        for i, line in enumerate(lines):
            para = frame.paragraphs[0] if i == 0 else frame.add_paragraph()
            para.text = f"• {line}"
            para.font.size, para.font.color.rgb = Pt(size), rgb(INK)
            para.space_after = Pt(6)

    def grid(
        slide: Any,
        x: float,
        y: float,
        w: float,
        headers: list[str],
        rows: list[list[str]],
        *,
        size: int = 12,
    ) -> None:
        shape = slide.shapes.add_table(
            len(rows) + 1,
            len(headers),
            Inches(x),
            Inches(y),
            Inches(w),
            Inches(0.4 * (len(rows) + 1)),
        )
        tbl = shape.table
        for c, head in enumerate(headers):
            cell = tbl.cell(0, c)
            cell.text = head
            cell.fill.solid()
            cell.fill.fore_color.rgb = rgb(BRAND)
            para = cell.text_frame.paragraphs[0]
            para.font.size, para.font.bold, para.font.color.rgb = Pt(size), True, rgb("FFFFFF")
        for r, values in enumerate(rows, 1):
            for c, value in enumerate(values):
                cell = tbl.cell(r, c)
                cell.text = value
                cell.fill.solid()
                cell.fill.fore_color.rgb = rgb("FFFFFF" if r % 2 else WASH)
                para = cell.text_frame.paragraphs[0]
                para.font.size, para.font.color.rgb = Pt(size), rgb(INK)

    # 1. Title
    slide = prs.slides.add_slide(blank)
    band = slide.shapes.add_shape(1, 0, 0, width, prs.slide_height)
    band.fill.solid()
    band.fill.fore_color.rgb = rgb(BRAND)
    band.line.fill.background()
    text(slide, 0.9, 2.3, 11.5, 0.5, brand.upper(), size=14, bold=True, colour="BFDBFE")
    text(slide, 0.9, 2.8, 11.5, 1.2, report.title, size=44, bold=True, colour="FFFFFF")
    text(
        slide,
        0.9,
        4.1,
        11.5,
        0.6,
        f"{period_range(report)} · {scope_text(report)}",
        size=20,
        colour="FFFFFF",
    )

    # 2. Headline numbers + summary
    slide = slide_with_title("Headline numbers")
    cards = kpi_cards(report)
    card_w = 2.85
    for i, (label, value, note, colour) in enumerate(cards):
        x = 0.6 + i * (card_w + 0.25)
        box = slide.shapes.add_shape(5, Inches(x), Inches(1.4), Inches(card_w), Inches(1.9))
        box.fill.solid()
        box.fill.fore_color.rgb = rgb(WASH)
        box.line.color.rgb = rgb(RULE)
        box.adjustments[0] = 0.08
        text(
            slide, x + 0.2, 1.55, card_w - 0.4, 0.4, label.upper(), size=11, bold=True, colour=MUTED
        )
        text(slide, x + 0.2, 1.95, card_w - 0.4, 0.8, value, size=36, bold=True, colour=colour)
        text(slide, x + 0.2, 2.75, card_w - 0.4, 0.5, note, size=11, colour=MUTED)
    bullets(slide, 0.6, 3.7, 12.0, 3.4, report.summary, size=18)

    # 3. Trend charts
    slide = slide_with_title("Trend")
    days = [p.day[5:] for p in report.trend]
    rate = CategoryChartData()
    rate.categories = days
    rate.add_series("Pass rate %", [p.pass_rate for p in report.trend])
    frame = slide.shapes.add_chart(
        XL_CHART_TYPE.LINE_MARKERS, Inches(0.5), Inches(1.4), Inches(6.1), Inches(5.4), rate
    )
    chart = frame.chart
    chart.font.size = Pt(12)
    chart.has_title = True
    chart.chart_title.text_frame.text = "Pass rate by day (%)"
    chart.has_legend = False
    known = [p.pass_rate for p in report.trend if p.pass_rate is not None]
    if known:
        chart.value_axis.minimum_scale = max(0.0, (min(known) // 5) * 5 - 5)
        chart.value_axis.maximum_scale = 100.0
        chart.value_axis.major_unit = 5.0
    line = chart.plots[0].series[0]
    line.format.line.color.rgb = rgb(BRAND)
    line.format.line.width = Emu(28575)

    runs = CategoryChartData()
    runs.categories = days
    runs.add_series("Green", [p.green for p in report.trend])
    runs.add_series("Red", [p.runs - p.green for p in report.trend])
    frame = slide.shapes.add_chart(
        XL_CHART_TYPE.COLUMN_STACKED, Inches(6.8), Inches(1.4), Inches(6.1), Inches(5.4), runs
    )
    chart = frame.chart
    chart.font.size = Pt(12)
    chart.has_title = True
    chart.chart_title.text_frame.text = "Runs by day"
    chart.has_legend = True
    chart.legend.position = XL_LEGEND_POSITION.BOTTOM
    chart.legend.include_in_layout = False
    for series, colour in zip(chart.plots[0].series, (GOOD, BAD), strict=True):
        series.format.fill.solid()
        series.format.fill.fore_color.rgb = rgb(colour)

    # 4. Cycles
    slide = slide_with_title("Cycles")
    if report.cycles:
        grid(
            slide,
            0.6,
            1.4,
            12.1,
            ["Cycle", "Runs", "Green", "Pass rate", "Change", "Failing", "New flaky", "Health"],
            [
                [
                    c.cycle,
                    str(c.runs),
                    str(c.green),
                    _pct(c.pass_rate),
                    _delta(c.pass_rate, c.prev_pass_rate) or "–",
                    str(c.failing),
                    str(c.newly_flaky),
                    "–" if c.health_score is None else str(c.health_score),
                ]
                for c in report.cycles[:12]
            ],
            size=15,
        )
        if len(report.cycles) > 12:
            text(
                slide,
                0.6,
                6.8,
                12,
                0.4,
                f"… and {len(report.cycles) - 12} more cycles",
                size=12,
                colour=MUTED,
            )
    else:
        text(slide, 0.6, 1.5, 12, 0.6, "No cycle ran in this period.", size=18, colour=MUTED)

    # 5. What needs attention
    slide = slide_with_title("What needs attention")
    rows = [
        [label, t.name, t.cycle, t.detail]
        for label, items in _lists(report)
        if label != "Fixed"
        for t in items
    ][:12]
    if rows:
        grid(slide, 0.6, 1.4, 12.1, ["", "Test", "Cycle", "Change"], rows, size=14)
    else:
        text(
            slide,
            0.6,
            1.5,
            12,
            0.6,
            "Nothing is failing, newly flaky or slower.",
            size=18,
            colour=GOOD,
        )

    # 6. Good news and next steps
    slide = slide_with_title("Fixed, and what to fix next")
    fixed = [f"{t.name} ({t.cycle})" for t in report.fixed] or ["No test was fixed this period."]
    text(slide, 0.6, 1.3, 5.8, 0.5, "Fixed", size=18, bold=True, colour=GOOD)
    bullets(slide, 0.6, 1.8, 5.8, 5.2, fixed[:10], size=16)
    nexts = [f"{f.title} · {f.cycle}: {f.detail}" for f in report.findings] or [
        "Nothing stands out: keep it green."
    ]
    text(slide, 6.8, 1.3, 6.0, 0.5, "What to fix first", size=18, bold=True, colour=BRAND)
    bullets(slide, 6.8, 1.8, 6.0, 5.2, nexts[:8], size=16)

    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()
